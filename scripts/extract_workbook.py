"""Ekstrak workbook wedding planner menjadi data JSON untuk aplikasi.

Script ini sengaja menyimpan data mentah dan data yang telah dinormalisasi. Data
mentah menjaga semua isi sel non-kosong, sementara data normalized dipakai UI.
"""

from __future__ import annotations

import json
import re
from collections import defaultdict
from datetime import date, datetime
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter


ROOT = Path(__file__).resolve().parents[1]
WORKBOOK = next((ROOT / "docs").glob("*.xlsx"))
OUTPUT = ROOT / "data" / "wedding-data.json"
ASSET_DIR = ROOT / "public" / "excel-assets"


def serial(value):
    if isinstance(value, (datetime, date)):
        return value.isoformat()
    return value


def text(value):
    return "" if value is None else str(value).strip()


def number(value):
    return float(value or 0) if isinstance(value, (int, float)) else 0


def resolved(ws, coordinate):
    """Ambil nilai sel, termasuk nilai induk untuk merged cell."""
    cell = ws[coordinate]
    if cell.value is not None:
        return cell.value
    for merged in ws.merged_cells.ranges:
        if coordinate in merged:
            return ws.cell(merged.min_row, merged.min_col).value
    return None


def slug(value):
    return re.sub(r"[^a-z0-9]+", "-", value.lower()).strip("-")


def raw_sheet(ws):
    cells = []
    for row in ws.iter_rows():
        for cell in row:
            if cell.value is not None:
                cells.append({"cell": cell.coordinate, "value": serial(cell.value)})
    return {"name": ws.title, "cells": cells}


def extract_images(workbook):
    ASSET_DIR.mkdir(parents=True, exist_ok=True)
    assets = []
    for ws in workbook.worksheets:
        for index, image in enumerate(ws._images, 1):
            extension = (getattr(image, "format", None) or "png").lower()
            if extension not in {"png", "jpeg", "jpg", "gif"}:
                extension = "png"
            filename = f"{slug(ws.title)}-{index}.{extension}"
            (ASSET_DIR / filename).write_bytes(image._data())
            anchor = getattr(image, "anchor", None)
            origin = getattr(anchor, "_from", None)
            assets.append({
                "sheet": ws.title,
                "src": f"/excel-assets/{filename}",
                "row": getattr(origin, "row", 0) + 1 if origin else None,
                "column": get_column_letter(getattr(origin, "col", 0) + 1) if origin else None,
                "width": round(getattr(image, "width", 0)),
                "height": round(getattr(image, "height", 0)),
            })
    return assets


def extract_mua(ws):
    vendors = {}
    for row in range(4, ws.max_row + 1):
        vendor_name = text(resolved(ws, f"C{row}"))
        if not vendor_name or vendor_name.lower() == "vendor":
            continue
        vendor = vendors.setdefault(vendor_name, {
            "id": f"mua-{slug(vendor_name)}", "name": vendor_name,
            "social": text(resolved(ws, f"H{row}")),
            "contact": text(resolved(ws, f"I{row}")),
            "availability": text(resolved(ws, f"J{row}")), "packages": []
        })
        package_name = text(resolved(ws, f"D{row}")) or "Paket"
        package_key = (package_name, number(resolved(ws, f"F{row}")), number(resolved(ws, f"G{row}")))
        package = next((item for item in vendor["packages"] if item["_key"] == package_key), None)
        if not package:
            package = {
                "_key": package_key, "name": package_name,
                "ownerPrice": package_key[1], "teamPrice": package_key[2], "includes": []
            }
            vendor["packages"].append(package)
        inclusion = text(ws[f"E{row}"].value)
        if inclusion and inclusion not in package["includes"]:
            package["includes"].append(inclusion)
        for key, column in (("social", "H"), ("contact", "I"), ("availability", "J")):
            if not vendor[key]:
                vendor[key] = text(resolved(ws, f"{column}{row}"))
    for vendor in vendors.values():
        for package in vendor["packages"]:
            package.pop("_key", None)
    return list(vendors.values())


def main():
    formulas = load_workbook(WORKBOOK, data_only=False)
    values = load_workbook(WORKBOOK, data_only=True)

    summary = values["Summary"]
    expenses_ws = values["Pengeluaran"]
    expenses = []
    for row in range(5, expenses_ws.max_row + 1):
        detail = text(expenses_ws[f"D{row}"].value)
        amount = number(expenses_ws[f"G{row}"].value)
        if not detail and not amount:
            continue
        expenses.append({
            "id": f"excel-expense-{row}", "date": serial(expenses_ws[f"C{row}"].value),
            "name": detail or "Pengeluaran", "category": text(expenses_ws[f"E{row}"].value) or "Lainnya",
            "paymentType": text(expenses_ws[f"F{row}"].value) or "Pembayaran",
            "amount": amount, "note": text(expenses_ws[f"H{row}"].value),
            "vendor": "", "status": "paid"
        })

    actual_by_category = defaultdict(float)
    for expense in expenses:
        actual_by_category[expense["category"]] += expense["amount"]

    categories = []
    for row in range(13, 30):
        name = text(summary[f"B{row}"].value)
        if name:
            categories.append({
                "id": slug(name), "name": name, "allocation": number(summary[f"C{row}"].value),
                "actual": actual_by_category[name], "completed": bool(summary[f"F{row}"].value)
            })

    checklist_ws = values["Checklist"]
    checklist = []
    for row in range(3, 50):
        task = text(checklist_ws[f"A{row}"].value)
        if task:
            checklist.append({
                "id": f"task-{row}", "task": task,
                "status": text(checklist_ws[f"B{row}"].value) or ("DONE" if checklist_ws[f"C{row}"].value else "BELUM"),
                "done": bool(checklist_ws[f"C{row}"].value), "vendor": text(checklist_ws[f"D{row}"].value),
                "dueDate": serial(checklist_ws[f"E{row}"].value), "notes": text(checklist_ws[f"G{row}"].value)
            })

    locations_ws = values["Prewedding"]
    locations = []
    for row in range(5, locations_ws.max_row + 1):
        name = text(locations_ws[f"B{row}"].value)
        price = locations_ws[f"C{row}"].value
        if name and isinstance(price, (int, float)):
            locations.append({"id": f"location-{row}", "name": name, "price": number(price), "note": text(locations_ws[f"D{row}"].value)})

    guest_ws = values["Daftar Undangan"]
    guest_blocks = [
        ("Pengantin Pria", "C", "D", "E"), ("Orang Tua Pria", "G", "H", "I"),
        ("Pengantin Wanita", "O", "P", "Q"), ("Orang Tua Wanita", "S", "T", "U")
    ]
    guests = []
    for owner, name_col, priority_col, count_col in guest_blocks:
        for row in range(6, 56):
            name = text(guest_ws[f"{name_col}{row}"].value)
            if name:
                guests.append({
                    "id": f"guest-{slug(owner)}-{row}", "name": name, "owner": owner,
                    "priority": text(guest_ws[f"{priority_col}{row}"].value) or "Prioritas",
                    "pax": int(number(guest_ws[f"{count_col}{row}"].value) or 1), "rsvp": "Menunggu"
                })

    souvenir_ws = values["Souvenir dan Seserahan"]
    souvenirs = []
    for row in range(5, 11):
        vals = [souvenir_ws.cell(row, col).value for col in range(2, 6)]
        if any(value is not None for value in vals):
            souvenirs.append({"id": f"souvenir-{row}", "vendor": text(vals[0]), "item": text(vals[1]), "price": number(vals[2]), "link": text(vals[3])})
    offerings = []
    for row in range(14, 40):
        vals = [souvenir_ws.cell(row, col).value for col in range(2, 7)]
        if any(value not in (None, 0, False) for value in vals):
            offerings.append({"id": f"offering-{row}", "box": text(resolved(souvenir_ws, f"B{row}")), "item": text(vals[1]), "vendor": text(vals[2]), "price": number(vals[3]), "done": bool(vals[4])})

    docs_ws = values["Dokumen"]
    documents = []
    for row in range(2, docs_ws.max_row + 1):
        requirement = text(docs_ws[f"B{row}"].value)
        if requirement and requirement != "Administrasi":
            documents.append({
                "id": f"document-{row}", "requirement": requirement,
                "note": text(docs_ws[f"I{row}"].value), "link": text(docs_ws[f"J{row}"].value), "done": False
            })

    mood_ws = values["Mood Board"]
    mood_notes = []
    for row in mood_ws.iter_rows():
        for cell in row:
            value = text(cell.value)
            if value:
                mood_notes.append({"cell": cell.coordinate, "text": value})

    assets = extract_images(formulas)
    output = {
        "meta": {"source": WORKBOOK.name, "version": 1},
        "project": {
            "couple": text(summary["B8"].value) or "Rama & Sita",
            "events": [
                {"id": "ngidih", "name": text(summary["H2"].value), "date": serial(summary["H3"].value)},
                {"id": "reception", "name": text(summary["I2"].value), "date": serial(summary["I3"].value)}
            ],
            "fundingSources": [
                {"id": f"fund-{row}", "name": text(summary[f"H{row}"].value), "amount": number(summary[f"I{row}"].value)}
                for row in range(17, 22) if summary[f"H{row}"].value
            ]
        },
        "categories": categories, "expenses": expenses, "checklist": checklist,
        "muaVendors": extract_mua(values["MUA"]), "preweddingLocations": locations,
        "guests": guests, "souvenirs": souvenirs, "offerings": offerings,
        "documents": documents, "moodNotes": mood_notes, "assets": assets,
        "rawSheets": [raw_sheet(sheet) for sheet in values.worksheets]
    }
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    OUTPUT.write_text(json.dumps(output, ensure_ascii=False, indent=2), encoding="utf-8")
    print(f"Created {OUTPUT.relative_to(ROOT)} ({OUTPUT.stat().st_size:,} bytes)")
    print(f"Extracted {len(assets)} image assets")
    print(json.dumps({key: len(value) for key, value in output.items() if isinstance(value, list)}, indent=2))


if __name__ == "__main__":
    main()
